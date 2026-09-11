from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "resume-and-interview.md"
OUTPUT = ROOT / "docs" / "AI应用开发实习面试材料-朱奕韬.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17324D"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
WHITE = "FFFFFF"
BLACK = "1F2933"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa, strict=True):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, italic: bool | None = None,
                 font: str = "Calibri", east_asia: str = "Microsoft YaHei") -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text: str) -> None:
    token_re = re.compile(r"(\*\*.+?\*\*|`.+?`|<https?://[^>]+>)")
    cursor = 0
    for match in token_re.finditer(text):
        if match.start() > cursor:
            set_run_font(paragraph.add_run(text[cursor:match.start()]), color=BLACK)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, bold=True, color=NAVY)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=9.5, color=DARK_BLUE, font="Consolas", east_asia="Microsoft YaHei")
            run.font.highlight_color = None
        else:
            url = token[1:-1]
            run = paragraph.add_run(url)
            set_run_font(run, color=BLUE)
            run.underline = True
        cursor = match.end()
    if cursor < len(text):
        set_run_font(paragraph.add_run(text[cursor:]), color=BLACK)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.22

    values = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (11.5, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in values.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.2


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    end = paragraph.add_run(" 页")
    set_run_font(end, size=9, color=MUTED)


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    header = section.header.paragraphs[0]
    header.text = "AI 应用开发实习｜面试准备手册"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, size=8.5, color=MUTED)
    add_page_field(section.footer.paragraphs[0])


def add_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(12)
    set_run_font(kicker.add_run("AI APPLICATION ENGINEERING"), size=10, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    set_run_font(title.add_run("AI 应用开发实习\n面试准备手册"), size=28, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    set_run_font(subtitle.add_run("智能出行 Agent · RAG 文档知识库 · 电动汽车充电需求研究"), size=12.5, color=MUTED)

    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [3120, 3120, 3120], indent_dxa=0)
    values = [("57", "高频问题"), ("3", "项目主题"), ("3 步", "技术回答框架")]
    for cell, (value, label) in zip(table.rows[0].cells, values, strict=True):
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(value), size=19, bold=True, color=BLUE)
        p.add_run("\n")
        set_run_font(p.add_run(label), size=9.5, color=MUTED)

    doc.add_paragraph().paragraph_format.space_after = Pt(28)
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(5)
    set_run_font(name.add_run("朱奕韬"), size=14, bold=True, color=NAVY)
    target = doc.add_paragraph()
    target.alignment = WD_ALIGN_PARAGRAPH.CENTER
    target.paragraph_format.space_after = Pt(4)
    set_run_font(target.add_run("求职方向：AI 应用开发 / 大模型应用实习"), size=10.5, color=MUTED)
    stamp = doc.add_paragraph()
    stamp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(stamp.add_run(f"更新日期：{date.today().isoformat()}"), size=9, color=MUTED)
    stamp.add_run().add_break(WD_BREAK.PAGE)


def add_static_toc(doc: Document) -> None:
    doc.add_heading("使用说明与内容导航", level=1)
    lead = doc.add_paragraph()
    add_inline(lead, "建议先熟练掌握 **30 秒自我介绍** 和 **为什么用—怎么用—有什么好处** 三段式框架，再按岗位重点复习对应项目。")

    items = [
        ("01", "自我介绍与 Agent 项目介绍", "建立开场表达"),
        ("02", "核心技术三问速答", "应对技术选型追问"),
        ("03", "智能出行 Agent：1–27", "工具编排、MCP、HITL 与工程化"),
        ("04", "RAG 知识库助手：28–44", "索引、检索、引用与拒答"),
        ("05", "充电需求研究：45–52", "LDA、XGBoost 与模型解释"),
        ("06", "综合问题：53–57", "项目取舍、问题定位与成长方向"),
    ]
    table = doc.add_table(rows=0, cols=3)
    set_table_geometry(table, [780, 4760, 3820])
    for idx, title, note in items:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell, top=130, bottom=130)
        set_cell_shading(cells[0], BLUE)
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p0.add_run(idx), size=10, bold=True, color=WHITE)
        set_run_font(cells[1].paragraphs[0].add_run(title), size=10.5, bold=True, color=NAVY)
        set_run_font(cells[2].paragraphs[0].add_run(note), size=9.5, color=MUTED)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_tech_block(doc: Document, columns: list[str]) -> None:
    technology, why, how, benefit = columns
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(5)
    set_run_font(p.add_run(technology), size=11.5, bold=True, color=BLUE)
    for label, value in (("为什么用", why), ("怎么用", how), ("带来的好处", benefit)):
        para = cell.add_paragraph()
        para.paragraph_format.space_after = Pt(3)
        set_run_font(para.add_run(f"{label}："), size=9.8, bold=True, color=NAVY)
        add_inline(para, value)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def render_markdown(doc: Document, text: str) -> None:
    lines = text.splitlines()
    i = 1  # skip source H1; the Word document has its own cover
    in_tech_tables = False
    pending_paragraph: list[str] = []

    def flush_paragraph() -> None:
        nonlocal pending_paragraph
        if not pending_paragraph:
            return
        paragraph = doc.add_paragraph()
        add_inline(paragraph, " ".join(part.strip() for part in pending_paragraph))
        pending_paragraph = []

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            flush_paragraph()
            i += 1
            continue

        if line == "## 简历项目描述（可直接粘贴）":
            flush_paragraph()
            # This Word file is an interview handbook, so skip the resume-only block.
            i += 1
            while i < len(lines) and lines[i].strip() != "## 简历数字的正确讲法":
                i += 1
            continue

        if line.startswith("## "):
            flush_paragraph()
            title = line[3:]
            doc.add_heading(title, level=1)
            in_tech_tables = title == "核心技术三问速答"
            i += 1
            continue
        if line.startswith("### "):
            flush_paragraph()
            doc.add_heading(line[4:], level=2)
            i += 1
            continue

        if in_tech_tables and line.startswith("| 技术 |"):
            flush_paragraph()
            i += 2  # header + separator
            while i < len(lines) and lines[i].strip().startswith("|"):
                columns = [part.strip() for part in lines[i].strip().strip("|").split("|")]
                if len(columns) == 4:
                    add_tech_block(doc, columns)
                i += 1
            continue

        if re.match(r"^\d+\. ", line):
            flush_paragraph()
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\. ", "", line))
            i += 1
            continue
        if line.startswith("- "):
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
            i += 1
            continue
        if line.startswith("> "):
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.right_indent = Inches(0.18)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, line[2:])
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
            i += 1
            continue

        pending_paragraph.append(line)
        i += 1

    flush_paragraph()


def keep_question_with_answer(doc: Document) -> None:
    paragraphs = doc.paragraphs
    for index, paragraph in enumerate(paragraphs[:-1]):
        if paragraph.style.name == "Heading 2" and re.match(r"^\d+\.", paragraph.text):
            paragraph.paragraph_format.keep_with_next = True
            paragraphs[index + 1].paragraph_format.keep_together = True


def build() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    add_cover(doc)
    add_static_toc(doc)
    render_markdown(doc, text)
    keep_question_with_answer(doc)

    core = doc.core_properties
    core.title = "AI 应用开发实习面试准备手册"
    core.subject = "智能出行 Agent、RAG 知识库与充电需求研究高频问答"
    core.author = "朱奕韬"
    core.keywords = "AI应用开发, Agent, RAG, MCP, LangGraph, Milvus, 面试"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
